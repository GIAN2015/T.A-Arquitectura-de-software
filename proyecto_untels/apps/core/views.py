from django.shortcuts import render, redirect, get_object_or_404
from django.contrib import messages
from django.http import HttpResponse
from django.utils import timezone
from apps.usuarios.services import identificar_usuario, autenticar_usuario, registrar_usuario
from apps.usuarios.models import Usuario
from apps.informes.models import Informe
from apps.informes.services import leer_archivo
from apps.reglamento.services import obtener_reglamento
from apps.observaciones.services import obtener_observaciones, validar_informe
from apps.observaciones.models import ObservacionGenerada


# ============================================
# AUTENTICACION
# ============================================

def login_view(request):
    if 'usuario_id' in request.session:
        if request.session.get('usuario_tipo') == 'docente':
            return redirect('panel_docente')
        return redirect('upload')

    if request.method == 'POST':
        codigo = request.POST.get('codigo', '').strip()
        password = request.POST.get('password', '').strip()
        if not codigo or not password:
            messages.error(request, 'Por favor ingresa tu codigo y contrasena.')
            return render(request, 'login.html')
        usuario = autenticar_usuario(codigo, password)
        if usuario:
            if usuario.tipo_usuario == 'docente':
                messages.error(request, 'Los docentes deben usar el portal de docentes.')
                return render(request, 'login.html')
            request.session['usuario_id'] = usuario.id
            request.session['usuario_codigo'] = usuario.codigo
            request.session['usuario_nombre'] = usuario.nombre
            request.session['usuario_tipo'] = usuario.tipo_usuario
            return redirect('upload')
        else:
            messages.error(request, 'Codigo o contrasena incorrectos.')
    return render(request, 'login.html')


def login_docente_view(request):
    if 'usuario_id' in request.session and request.session.get('usuario_tipo') == 'docente':
        return redirect('panel_docente')

    if request.method == 'POST':
        codigo = request.POST.get('codigo', '').strip()
        password = request.POST.get('password', '').strip()
        if not codigo or not password:
            messages.error(request, 'Por favor ingresa tu codigo y contrasena.')
            return render(request, 'login_docente.html')
        usuario = autenticar_usuario(codigo, password)
        if usuario:
            if usuario.tipo_usuario != 'docente':
                messages.error(request, 'Esta area es solo para docentes.')
                return render(request, 'login_docente.html')
            request.session['usuario_id'] = usuario.id
            request.session['usuario_codigo'] = usuario.codigo
            request.session['usuario_nombre'] = usuario.nombre
            request.session['usuario_tipo'] = usuario.tipo_usuario
            return redirect('panel_docente')
        else:
            messages.error(request, 'Codigo o contrasena incorrectos.')
    return render(request, 'login_docente.html')


def registro_view(request):
    if request.method == 'POST':
        codigo = request.POST.get('codigo', '').strip()
        nombre = request.POST.get('nombre', '').strip()
        password = request.POST.get('password', '').strip()
        password_confirm = request.POST.get('password_confirm', '').strip()
        tipo_usuario = request.POST.get('tipo_usuario', 'estudiante')

        if tipo_usuario == 'docente':
            messages.error(request, 'Los docentes son registrados por la administracion.')
            return render(request, 'registro.html')
        if not all([codigo, nombre, password, password_confirm]):
            messages.error(request, 'Todos los campos son obligatorios.')
            return render(request, 'registro.html')
        if password != password_confirm:
            messages.error(request, 'Las contrasenas no coinciden.')
            return render(request, 'registro.html')
        if len(password) < 6:
            messages.error(request, 'La contrasena debe tener al menos 6 caracteres.')
            return render(request, 'registro.html')

        usuario = registrar_usuario(codigo, nombre, password, tipo_usuario)
        if usuario:
            messages.success(request, 'Registro exitoso. Ahora puedes iniciar sesion.')
            return redirect('login')
        else:
            messages.error(request, 'El codigo ya esta registrado.')
    return render(request, 'registro.html')


def logout_view(request):
    tipo = request.session.get('usuario_tipo')
    request.session.flush()
    if tipo == 'docente':
        return redirect('login_docente')
    return redirect('login')


# ============================================
# VISTAS DE ESTUDIANTES
# ============================================

def _verificar_estudiante(request):
    if 'usuario_id' not in request.session:
        return False
    return request.session.get('usuario_tipo') != 'docente'


def upload_view(request):
    if not _verificar_estudiante(request):
        return redirect('login')

    context = {
        'codigo': request.session.get('usuario_codigo'),
        'nombre': request.session.get('usuario_nombre'),
        'tipo': request.session.get('usuario_tipo'),
    }

    if request.method == 'POST':
        archivo = request.FILES.get('informe')
        if not archivo:
            messages.error(request, 'Por favor selecciona un archivo .docx')
            return render(request, 'upload_report.html', context)
        if not archivo.name.endswith('.docx'):
            messages.error(request, 'Solo se permiten archivos .docx')
            return render(request, 'upload_report.html', context)

        informe = None
        try:
            usuario = Usuario.objects.get(id=request.session['usuario_id'])
            contenido = leer_archivo(archivo)

            informe_previo = Informe.objects.filter(
                usuario=usuario,
                estado=Informe.ESTADO_RECHAZADO
            ).order_by('-fecha_registro').first()

            version = 1
            if informe_previo:
                version = informe_previo.version + 1

            informe = Informe.objects.create(
                usuario=usuario,
                nombre_archivo=archivo.name,
                contenido=contenido,
                estado=Informe.ESTADO_ENVIADO,
                version=version,
                informe_anterior=informe_previo
            )

            informe.estado = Informe.ESTADO_VALIDANDO
            informe.save()

            reglamento = obtener_reglamento()
            observaciones_banco = obtener_observaciones()
            resultado = validar_informe(contenido, reglamento, observaciones_banco)

            for obs in resultado:
                ObservacionGenerada.objects.create(
                    informe=informe,
                    seccion=obs.get('descripcion', ''),
                    observacion=obs.get('observacion', ''),
                    sustento=obs.get('sustento', ''),
                    estado_conformidad=obs.get('estado', 'Observado'),
                )

            informe.estado = Informe.ESTADO_OBSERVADO
            informe.save()
            return redirect('resultado', informe_id=informe.id)

        except Exception as e:
            if informe:
                informe.estado = Informe.ESTADO_OBSERVADO
                informe.save()
            messages.error(request, f'Error al procesar el informe: {str(e)}')
            return render(request, 'upload_report.html', context)

    return render(request, 'upload_report.html', context)


def resultado_view(request, informe_id):
    if 'usuario_id' not in request.session:
        return redirect('login')

    informe = get_object_or_404(Informe, id=informe_id)
    usuario_id = request.session.get('usuario_id')
    tipo = request.session.get('usuario_tipo')

    if tipo != 'docente' and informe.usuario_id != usuario_id:
        messages.error(request, 'No tienes permiso para ver este informe.')
        return redirect('historial')

    observaciones = ObservacionGenerada.objects.filter(informe=informe)
    context = {
        'informe': informe,
        'observaciones': observaciones,
        'usuario': informe.usuario,
        'es_docente': tipo == 'docente',
    }
    return render(request, 'validation_result.html', context)


def historial_view(request):
    if not _verificar_estudiante(request):
        return redirect('login')

    usuario = Usuario.objects.get(id=request.session['usuario_id'])
    informes = Informe.objects.filter(usuario=usuario).order_by('-fecha_registro')

    context = {
        'codigo': request.session.get('usuario_codigo'),
        'nombre': request.session.get('usuario_nombre'),
        'tipo': request.session.get('usuario_tipo'),
        'informes': informes,
    }
    return render(request, 'historial.html', context)


# ============================================
# EXPORTAR EXCEL (nuestra mejora)
# ============================================

def exportar_excel(request, informe_id):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

    informe = get_object_or_404(Informe, id=informe_id)
    observaciones = ObservacionGenerada.objects.filter(informe=informe)

    wb = Workbook()
    ws = wb.active
    ws.title = "Observaciones"

    ws.merge_cells('A1:F1')
    ws['A1'] = f"INFORME DE OBSERVACIONES - {informe.nombre_archivo}"
    ws['A1'].font = Font(bold=True, size=14, color="FFFFFF")
    ws['A1'].fill = PatternFill(start_color="1A3A6B", end_color="1A3A6B", fill_type="solid")
    ws['A1'].alignment = Alignment(horizontal="center", vertical="center")

    ws.merge_cells('A2:F2')
    ws['A2'] = f"Alumno: {informe.usuario.nombre} | Codigo: {informe.usuario.codigo}"
    ws['A2'].font = Font(size=11)
    ws['A2'].alignment = Alignment(horizontal="center")

    headers = ['ITEM', 'DESCRIPCION', 'OBSERVACION', 'SUSTENTO', 'ESTADO', 'DOCENTE']
    header_fill = PatternFill(start_color="1A3A6B", end_color="1A3A6B", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF", size=11)
    thin_border = Border(
        left=Side(style='thin'), right=Side(style='thin'),
        top=Side(style='thin'), bottom=Side(style='thin'),
    )

    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=4, column=col, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = thin_border

    conforme_fill = PatternFill(start_color="D4EDDA", end_color="D4EDDA", fill_type="solid")
    observado_fill = PatternFill(start_color="FFF3CD", end_color="FFF3CD", fill_type="solid")

    for i, obs in enumerate(observaciones, 1):
        row = i + 4
        fill = conforme_fill if obs.estado_conformidad == 'Conforme' else observado_fill
        values = [i, obs.seccion, obs.observacion, obs.sustento, obs.estado_conformidad, obs.comentario_docente]
        for col, val in enumerate(values, 1):
            cell = ws.cell(row=row, column=col, value=val)
            cell.border = thin_border
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            if col == 5:
                cell.fill = fill

    ws.column_dimensions['A'].width = 6
    ws.column_dimensions['B'].width = 35
    ws.column_dimensions['C'].width = 45
    ws.column_dimensions['D'].width = 25
    ws.column_dimensions['E'].width = 14
    ws.column_dimensions['F'].width = 20

    response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    filename = f"observaciones_{informe.usuario.codigo}_{informe.id}.xlsx"
    response['Content-Disposition'] = f'attachment; filename="{filename}"'
    wb.save(response)
    return response


# ============================================
# VISTAS DE DOCENTES
# ============================================

def _verificar_docente(request):
    if 'usuario_id' not in request.session:
        return False
    return request.session.get('usuario_tipo') == 'docente'


def panel_docente_view(request):
    if not _verificar_docente(request):
        return redirect('login_docente')

    filtro_estado = request.GET.get('estado', 'pendientes')

    if filtro_estado == 'pendientes':
        informes = Informe.objects.filter(
            estado__in=[Informe.ESTADO_OBSERVADO, Informe.ESTADO_EN_REVISION_DOCENTE]
        ).select_related('usuario').order_by('-fecha_registro')
    elif filtro_estado == 'aprobados':
        informes = Informe.objects.filter(
            estado=Informe.ESTADO_APROBADO
        ).select_related('usuario').order_by('-fecha_revision_docente')
    elif filtro_estado == 'rechazados':
        informes = Informe.objects.filter(
            estado=Informe.ESTADO_RECHAZADO
        ).select_related('usuario').order_by('-fecha_revision_docente')
    else:
        informes = Informe.objects.all().select_related('usuario').order_by('-fecha_registro')

    total_pendientes = Informe.objects.filter(
        estado__in=[Informe.ESTADO_OBSERVADO, Informe.ESTADO_EN_REVISION_DOCENTE]
    ).count()
    total_aprobados = Informe.objects.filter(estado=Informe.ESTADO_APROBADO).count()
    total_rechazados = Informe.objects.filter(estado=Informe.ESTADO_RECHAZADO).count()
    total_informes = Informe.objects.count()

    context = {
        'codigo': request.session.get('usuario_codigo'),
        'nombre': request.session.get('usuario_nombre'),
        'informes': informes,
        'filtro_estado': filtro_estado,
        'total_pendientes': total_pendientes,
        'total_aprobados': total_aprobados,
        'total_rechazados': total_rechazados,
        'total_informes': total_informes,
    }
    return render(request, 'panel_docente.html', context)


def revisar_informe_view(request, informe_id):
    if not _verificar_docente(request):
        return redirect('login_docente')

    informe = get_object_or_404(Informe, id=informe_id)

    if informe.estado == Informe.ESTADO_OBSERVADO:
        informe.estado = Informe.ESTADO_EN_REVISION_DOCENTE
        informe.save()

    observaciones = informe.observaciones.all().order_by('id')

    if request.method == 'POST':
        accion = request.POST.get('accion')
        comentario_general = request.POST.get('comentario_general', '')

        for obs in observaciones:
            obs_comentario = request.POST.get(f'obs_comentario_{obs.id}', '')
            obs.comentario_docente = obs_comentario
            obs.fecha_revision = timezone.now()
            obs.save()

        docente = Usuario.objects.get(id=request.session.get('usuario_id'))
        informe.docente_revisor = docente
        informe.comentario_docente = comentario_general
        informe.fecha_revision_docente = timezone.now()

        if accion == 'aprobar':
            informe.estado = Informe.ESTADO_APROBADO
            messages.success(request, f'Informe "{informe.nombre_archivo}" aprobado.')
        elif accion == 'rechazar':
            informe.estado = Informe.ESTADO_RECHAZADO
            messages.warning(request, f'Informe "{informe.nombre_archivo}" rechazado. El alumno debe corregir.')

        informe.save()
        return redirect('panel_docente')

    context = {
        'nombre': request.session.get('usuario_nombre'),
        'informe': informe,
        'observaciones': observaciones,
    }
    return render(request, 'revisar_informe.html', context)
